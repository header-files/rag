"""Word (.docx) parser."""

from __future__ import annotations

import io
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from rag.errors import RagError
from rag.parsers.base import ParsedDocument


class DocxParser:
    suffixes = (".docx",)

    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        try:
            document = Document(io.BytesIO(data))
        except (PackageNotFoundError, ValueError, BadZipFile) as exc:
            raise RagError(
                message=f"Failed to parse docx '{filename}': {exc}",
                type="unsupported_file",
            ) from exc

        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = "\t".join(cell for cell in cells if cell)
                if line:
                    parts.append(line)
        return ParsedDocument(text="\n".join(parts))
